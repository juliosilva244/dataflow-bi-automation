from __future__ import annotations

import csv
import io
import zipfile
from pathlib import PurePosixPath

import pandas as pd
from openpyxl import load_workbook

from services.errors import UploadValidationError
from services.settings import SecuritySettings, get_security_settings

XLSX_REQUIRED_MEMBERS = {"[Content_Types].xml", "xl/workbook.xml"}


def read_limited_bytes(upload, settings: SecuritySettings) -> bytes:
    try:
        upload.seek(0)
        content = upload.read(settings.max_upload_bytes + 1)
        upload.seek(0)
    except (AttributeError, OSError, ValueError) as error:
        raise UploadValidationError(
            "Não foi possível ler o arquivo enviado.", "UPLOAD_UNREADABLE"
        ) from error

    if len(content) > settings.max_upload_bytes:
        limit_mb = settings.max_upload_bytes // (1024 * 1024)
        raise UploadValidationError(
            f"O arquivo excede o limite configurado de {limit_mb} MB.",
            "UPLOAD_TOO_LARGE",
        )
    if not content:
        raise UploadValidationError("O arquivo enviado está vazio.", "UPLOAD_EMPTY")
    return content


def _decode_csv(content: bytes) -> str:
    if b"\x00" in content:
        raise UploadValidationError(
            "O conteúdo não corresponde a um CSV de texto válido.", "CSV_BINARY_CONTENT"
        )
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return content.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise UploadValidationError(
        "Não foi possível identificar uma codificação CSV segura.", "CSV_ENCODING"
    )


def read_validated_csv(content: bytes, settings: SecuritySettings) -> pd.DataFrame:
    if content.startswith(b"PK\x03\x04"):
        raise UploadValidationError(
            "O conteúdo do arquivo não corresponde à extensão CSV.", "CSV_SIGNATURE_MISMATCH"
        )

    text = _decode_csv(content)
    sample = text[:16_384]
    try:
        delimiter = csv.Sniffer().sniff(sample, delimiters=",;\t|").delimiter
    except csv.Error:
        first_line = sample.splitlines()[0] if sample.splitlines() else ""
        delimiter = max(",;\t|", key=first_line.count)
        if first_line.count(delimiter) == 0:
            raise UploadValidationError(
                "A estrutura do CSV não pôde ser reconhecida.", "CSV_STRUCTURE"
            ) from None

    reader = csv.reader(io.StringIO(text, newline=""), delimiter=delimiter)
    try:
        header = next(reader)
    except (StopIteration, csv.Error) as error:
        raise UploadValidationError("O CSV não possui cabeçalho válido.", "CSV_HEADER") from error

    if not header or not any(str(value).strip() for value in header):
        raise UploadValidationError("O CSV não possui cabeçalho válido.", "CSV_HEADER")
    if len(header) > settings.max_columns:
        raise UploadValidationError(
            f"O arquivo excede o limite de {settings.max_columns} colunas.",
            "UPLOAD_TOO_MANY_COLUMNS",
        )

    row_count = 0
    try:
        for row in reader:
            if not row or not any(str(value).strip() for value in row):
                continue
            row_count += 1
            if len(row) != len(header):
                raise UploadValidationError(
                    "O CSV possui linhas com quantidade inconsistente de colunas.",
                    "CSV_INCONSISTENT_COLUMNS",
                )
            if row_count > settings.max_rows:
                raise UploadValidationError(
                    f"O arquivo excede o limite de {settings.max_rows:,} linhas.",
                    "UPLOAD_TOO_MANY_ROWS",
                )
    except csv.Error as error:
        raise UploadValidationError("O CSV possui estrutura inválida.", "CSV_STRUCTURE") from error

    if row_count == 0:
        raise UploadValidationError("O CSV não possui registros de dados.", "CSV_NO_DATA")

    return pd.read_csv(io.StringIO(text), sep=delimiter)


def _validate_zip(content: bytes, settings: SecuritySettings) -> None:
    if not zipfile.is_zipfile(io.BytesIO(content)):
        raise UploadValidationError(
            "O conteúdo do arquivo não corresponde a uma planilha XLSX válida.",
            "XLSX_SIGNATURE_MISMATCH",
        )

    with zipfile.ZipFile(io.BytesIO(content)) as archive:
        members = archive.infolist()
        names = {item.filename for item in members}
        if not XLSX_REQUIRED_MEMBERS.issubset(names):
            raise UploadValidationError("A estrutura interna do XLSX é inválida.", "XLSX_STRUCTURE")
        if len(members) > settings.max_xlsx_entries:
            raise UploadValidationError("A planilha possui estrutura interna excessiva.", "XLSX_ENTRIES")

        uncompressed_total = 0
        for item in members:
            path = PurePosixPath(item.filename)
            if path.is_absolute() or ".." in path.parts or item.flag_bits & 0x1:
                raise UploadValidationError(
                    "A planilha contém uma estrutura interna não permitida.",
                    "XLSX_UNSAFE_STRUCTURE",
                )
            uncompressed_total += item.file_size
            if uncompressed_total > settings.max_xlsx_uncompressed_bytes:
                raise UploadValidationError(
                    "O conteúdo descompactado da planilha excede o limite seguro.",
                    "XLSX_UNCOMPRESSED_TOO_LARGE",
                )
            if item.file_size and item.compress_size == 0:
                raise UploadValidationError("A planilha possui compressão inválida.", "XLSX_ZIP_BOMB")
            if item.compress_size and item.file_size / item.compress_size > settings.max_zip_compression_ratio:
                raise UploadValidationError(
                    "A taxa de compressão da planilha excede o limite seguro.",
                    "XLSX_ZIP_BOMB",
                )


def read_validated_xlsx(content: bytes, settings: SecuritySettings) -> pd.DataFrame:
    _validate_zip(content, settings)
    try:
        workbook = load_workbook(
            io.BytesIO(content), read_only=True, data_only=True, keep_links=False
        )
        worksheet = workbook.active
        rows = worksheet.iter_rows(values_only=True)
        header_row = next(rows, None)
        if header_row is None:
            raise UploadValidationError("A planilha não possui cabeçalho.", "XLSX_HEADER")
        header = list(header_row)
        while header and header[-1] is None:
            header.pop()
        if not header or not any(str(value).strip() for value in header if value is not None):
            raise UploadValidationError("A planilha não possui cabeçalho válido.", "XLSX_HEADER")
        if len(header) > settings.max_columns:
            raise UploadValidationError(
                f"O arquivo excede o limite de {settings.max_columns} colunas.",
                "UPLOAD_TOO_MANY_COLUMNS",
            )

        data = []
        for row in rows:
            values = list(row[: len(header)])
            if not any(value is not None and str(value).strip() for value in values):
                continue
            data.append(values)
            if len(data) > settings.max_rows:
                raise UploadValidationError(
                    f"O arquivo excede o limite de {settings.max_rows:,} linhas.",
                    "UPLOAD_TOO_MANY_ROWS",
                )
        if not data:
            raise UploadValidationError("A planilha não possui registros de dados.", "XLSX_NO_DATA")
        return pd.DataFrame(data, columns=header)
    except UploadValidationError:
        raise
    except Exception as error:
        raise UploadValidationError(
            "Não foi possível validar a estrutura da planilha XLSX.", "XLSX_PARSE"
        ) from error
    finally:
        if "workbook" in locals():
            workbook.close()


def read_validated_upload(upload, settings: SecuritySettings | None = None) -> pd.DataFrame:
    settings = settings or get_security_settings()
    content = read_limited_bytes(upload, settings)
    filename = str(getattr(upload, "name", "")).casefold()

    return read_validated_content(content, filename, settings)


def read_validated_content(
    content: bytes,
    filename: str,
    settings: SecuritySettings | None = None,
) -> pd.DataFrame:
    settings = settings or get_security_settings()
    filename = filename.casefold()

    if filename.endswith(".csv"):
        return read_validated_csv(content, settings)
    if filename.endswith(".xlsx"):
        return read_validated_xlsx(content, settings)
    raise UploadValidationError(
        "Formato não suportado. Envie somente arquivos CSV ou XLSX.", "UPLOAD_EXTENSION"
    )

