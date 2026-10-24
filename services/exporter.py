from io import BytesIO

import pandas as pd
from openpyxl.styles import (
    Alignment,
    Border,
    Font,
    PatternFill,
    Side,
)
from openpyxl.utils import get_column_letter

from services.insight_engine import gerar_insights_executivos
from services.security import neutralize_excel_formula
from services.settings import get_performance_settings

PRIMARY_BG = "0F172A"
SECONDARY_BG = "111827"
ACCENT = "38BDF8"
SUCCESS = "22C55E"
WARNING = "F59E0B"
TEXT_LIGHT = "F8FAFC"


def create_fill(color):
    return PatternFill(
        start_color=color,
        end_color=color,
        fill_type="solid",
    )


def create_font(
    size=11,
    bold=False,
    color=TEXT_LIGHT,
):
    return Font(
        name="Calibri",
        size=size,
        bold=bold,
        color=color,
    )


def create_border():
    side = Side(
        style="thin",
        color="334155",
    )

    return Border(
        left=side,
        right=side,
        top=side,
        bottom=side,
    )


def apply_header_style(ws, row=1):
    fill = create_fill(PRIMARY_BG)
    font = create_font(
        size=12,
        bold=True,
    )

    border = create_border()

    for cell in ws[row]:
        cell.fill = fill
        cell.font = font
        cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )
        cell.border = border


def autofit_columns(ws, sample_rows=None):
    sample_rows = sample_rows or get_performance_settings().excel_width_sample_rows
    for column_cells in ws.iter_cols(
        min_row=1,
        max_row=min(ws.max_row, sample_rows + 4),
    ):
        max_length = 0

        column_letter = get_column_letter(
            column_cells[0].column
        )

        for cell in column_cells:
            try:
                value_length = len(str(cell.value))
                max_length = max(
                    max_length,
                    value_length,
                )
            except Exception:
                pass

        adjusted_width = min(
            max_length + 5,
            42,
        )

        ws.column_dimensions[
            column_letter
        ].width = adjusted_width


def finalize_sheet(ws):
    apply_header_style(ws)
    autofit_columns(ws)
    ws.freeze_panes = "A2"


def create_title(
    ws,
    title,
    subtitle=None,
):
    ws.merge_cells("A1:F1")

    title_cell = ws["A1"]

    title_cell.value = title

    title_cell.font = create_font(
        size=18,
        bold=True,
    )

    title_cell.fill = create_fill(
        PRIMARY_BG
    )

    title_cell.alignment = Alignment(
        horizontal="center",
        vertical="center",
    )

    ws.row_dimensions[1].height = 32

    if subtitle:
        ws.merge_cells("A2:F2")

        subtitle_cell = ws["A2"]

        subtitle_cell.value = subtitle

        subtitle_cell.font = create_font(
            size=11,
            color="CBD5E1",
        )

        subtitle_cell.fill = create_fill(
            SECONDARY_BG
        )

        subtitle_cell.alignment = Alignment(
            horizontal="center",
            vertical="center",
        )

        ws.row_dimensions[2].height = 24


def build_executive_summary(metricas):
    return pd.DataFrame(
        {
            "Indicador": [
                "Faturamento Total",
                "Ticket Médio",
                "Pedidos",
                "Quantidade Vendida",
                "Lojas Ativas",
                "Melhor Loja",
                "Produto Campeão",
                "Participação Produto Top (%)",
                "Crescimento Faturamento (%)",
                "Crescimento Ticket Médio (%)",
            ],
            "Valor": [
                metricas.get(
                    "faturamento_total",
                    0,
                ),
                metricas.get(
                    "ticket_medio",
                    0,
                ),
                metricas.get(
                    "total_pedidos",
                    0,
                ),
                metricas.get(
                    "quantidade_total",
                    0,
                ),
                metricas.get(
                    "lojas_ativas",
                    0,
                ),
                metricas.get(
                    "melhor_loja",
                    "N/A",
                ),
                metricas.get(
                    "produto_top",
                    "N/A",
                ),
                round(
                    metricas.get(
                        "participacao_produto_top",
                        0,
                    ),
                    2,
                ),
                round(
                    metricas.get(
                        "crescimento_faturamento",
                        0,
                    ),
                    2,
                ),
                round(
                    metricas.get(
                        "crescimento_ticket",
                        0,
                    ),
                    2,
                ),
            ],
        }
    )


def build_insights_dataframe(metricas):
    insights = gerar_insights_executivos(
        metricas
    )

    rows = []

    for item in insights:
        rows.append(
            {
                "Tipo": item.get(
                    "badge",
                    "Insight",
                ),
                "Título": item.get(
                    "title",
                    "",
                ),
                "Descrição": item.get(
                    "text",
                    "",
                ),
            }
        )

    return pd.DataFrame(rows)


def build_quality_dataframe(analytics_state):
    quality = analytics_state.get("quality", {})
    rows = [
        {"Indicador": "Registros recebidos", "Valor": quality.get("registros_recebidos", 0)},
        {"Indicador": "Registros válidos", "Valor": quality.get("registros_validos", 0)},
        {"Indicador": "Registros descartados", "Valor": quality.get("registros_descartados", 0)},
        {"Indicador": "Duplicados removidos", "Valor": quality.get("registros_duplicados", 0)},
        {"Indicador": "Score do schema (%)", "Valor": quality.get("score_schema", 0)},
    ]
    for reason, count in quality.get("motivos_rejeicao", {}).items():
        rows.append({"Indicador": f"Rejeição: {reason}", "Valor": count})
    return pd.DataFrame(rows)


def safe_groupby(
    df,
    group_col,
    value_col,
    sort_desc=True,
):
    if (
        group_col not in df.columns
        or value_col not in df.columns
    ):
        return pd.DataFrame()

    grouped = (
        df.groupby(group_col, as_index=False)[
            value_col
        ]
        .sum()
        .sort_values(
            value_col,
            ascending=not sort_desc,
        )
    )

    return grouped


def add_section_label(
    ws,
    row,
    text,
):
    cell = ws.cell(
        row=row,
        column=1,
    )

    cell.value = text

    cell.font = create_font(
        size=14,
        bold=True,
    )

    cell.fill = create_fill(
        ACCENT
    )

    cell.alignment = Alignment(
        horizontal="left",
        vertical="center",
    )


def _column_number_format(header):
    name = str(header or "").casefold()
    if any(term in name for term in ("faturamento", "receita", "valor", "ticket", "lucro", "preço")):
        return 'R$ #,##0.00'
    if any(term in name for term in ("percent", "participação", "crescimento", "%")):
        return '0.00"%"'
    if any(term in name for term in ("quantidade", "pedidos", "registros", "lojas", "ocorrências")):
        return '#,##0'
    if any(term in name for term in ("data", "período")):
        return "dd/mm/yyyy"
    return None


def apply_semantic_formats(ws, header_row=4):
    """Format only columns with known semantics, avoiding a full style matrix."""
    for column in range(1, ws.max_column + 1):
        number_format = _column_number_format(ws.cell(header_row, column).value)
        if number_format is None:
            continue
        for row in range(header_row + 1, ws.max_row + 1):
            ws.cell(row, column).number_format = number_format


def apply_indicator_formats(ws, header_row=4):
    """Handle mixed-type KPI tables whose value semantics are stored per row."""
    for row in range(header_row + 1, ws.max_row + 1):
        indicator = str(ws.cell(row, 1).value or "").casefold()
        value_cell = ws.cell(row, 2)
        if not isinstance(value_cell.value, (int, float)):
            continue
        if any(term in indicator for term in ("faturamento", "ticket", "receita", "lucro", "preço")):
            value_cell.number_format = 'R$ #,##0.00'
        elif any(term in indicator for term in ("%", "percent", "participação", "crescimento", "score")):
            value_cell.number_format = '0.00"%"'
        else:
            value_cell.number_format = '#,##0'


def sanitize_dataframe_for_excel(df):
    safe_df = df.copy()
    for column in safe_df.columns:
        safe_df[column] = safe_df[column].map(neutralize_excel_formula)
    safe_df.columns = [neutralize_excel_formula(str(column)) for column in safe_df.columns]
    return safe_df


def gerar_excel(
    df,
    analytics_state,
):
    output = BytesIO()

    metricas = analytics_state

    charts = analytics_state.get(
        "charts",
        {},
    )

    revenue_by_store = charts.get(
        "revenue_by_store",
        pd.DataFrame(),
    )

    top_products = charts.get(
        "top_products",
        pd.DataFrame(),
    )

    revenue_by_month = charts.get(
        "revenue_by_month",
        pd.DataFrame(),
    )

    executive_summary = sanitize_dataframe_for_excel(
        build_executive_summary(
            metricas
        )
    )

    insights_df = sanitize_dataframe_for_excel(
        build_insights_dataframe(
            metricas
        )
    )
    quality_df = sanitize_dataframe_for_excel(build_quality_dataframe(analytics_state))

    revenue_by_store = sanitize_dataframe_for_excel(revenue_by_store)
    top_products = sanitize_dataframe_for_excel(top_products)
    revenue_by_month = sanitize_dataframe_for_excel(revenue_by_month)
    df = sanitize_dataframe_for_excel(df)

    with pd.ExcelWriter(
        output,
        engine="openpyxl",
    ) as writer:

        executive_summary.to_excel(
            writer,
            sheet_name="Resumo_Executivo",
            index=False,
            startrow=3,
        )

        insights_df.to_excel(
            writer,
            sheet_name="Insights_IA",
            index=False,
            startrow=3,
        )

        quality_df.to_excel(
            writer,
            sheet_name="Qualidade_Dados",
            index=False,
            startrow=3,
        )

        revenue_by_store.to_excel(
            writer,
            sheet_name="Performance_Lojas",
            index=False,
            startrow=3,
        )

        top_products.to_excel(
            writer,
            sheet_name="Top_Produtos",
            index=False,
            startrow=3,
        )

        revenue_by_month.to_excel(
            writer,
            sheet_name="Receita_Mensal",
            index=False,
            startrow=3,
        )

        df.to_excel(
            writer,
            sheet_name="Dados_Filtrados",
            index=False,
            startrow=3,
        )

        workbook = writer.book

        summary_ws = workbook[
            "Resumo_Executivo"
        ]

        create_title(
            summary_ws,
            "DATAFLOW BI AUTOMATION",
            "Relatório Executivo Analítico",
        )

        insights_ws = workbook[
            "Insights_IA"
        ]

        create_title(
            insights_ws,
            "INTELIGÊNCIA EXECUTIVA",
            "Insights automáticos da operação",
        )

        quality_ws = workbook["Qualidade_Dados"]
        create_title(
            quality_ws,
            "QUALIDADE DOS DADOS",
            "Registros, perdas e motivos verificáveis da normalização",
        )

        stores_ws = workbook[
            "Performance_Lojas"
        ]

        create_title(
            stores_ws,
            "PERFORMANCE POR LOJA",
            "Ranking operacional consolidado",
        )

        products_ws = workbook[
            "Top_Produtos"
        ]

        create_title(
            products_ws,
            "TOP PRODUTOS",
            "Produtos líderes por faturamento",
        )

        revenue_ws = workbook[
            "Receita_Mensal"
        ]

        create_title(
            revenue_ws,
            "EVOLUÇÃO DE RECEITA",
            "Performance temporal consolidada",
        )

        data_ws = workbook[
            "Dados_Filtrados"
        ]

        create_title(
            data_ws,
            "BASE ANALÍTICA",
            "Dados filtrados exportados do dashboard",
        )

        for sheet_name in workbook.sheetnames:
            ws = workbook[
                sheet_name
            ]

            apply_header_style(
                ws,
                row=4,
            )

            autofit_columns(ws)

            apply_semantic_formats(ws, header_row=4)

            ws.freeze_panes = "A5"

        apply_indicator_formats(summary_ws, header_row=4)
        apply_indicator_formats(quality_ws, header_row=4)

        summary_ws.sheet_view.showGridLines = False
        insights_ws.sheet_view.showGridLines = False
        quality_ws.sheet_view.showGridLines = False
        stores_ws.sheet_view.showGridLines = False
        products_ws.sheet_view.showGridLines = False
        revenue_ws.sheet_view.showGridLines = False
        data_ws.sheet_view.showGridLines = False

    output.seek(0)

    return output
