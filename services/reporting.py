from __future__ import annotations

import re
import unicodedata
import zipfile
from dataclasses import dataclass
from datetime import datetime
from decimal import Decimal, InvalidOperation
from io import BytesIO, StringIO
from typing import Any

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter
from reportlab.lib import colors
from reportlab.lib.enums import TA_CENTER
from reportlab.lib.pagesizes import A4, landscape
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import mm
from reportlab.platypus import (
    PageBreak,
    Paragraph,
    SimpleDocTemplate,
    Spacer,
    Table,
    TableStyle,
)

from services.data_quality import audit_data_quality, findings_dataframe
from services.metrics_engine import construir_analytics_state, exact_decimal_sum
from services.security import neutralize_excel_formula


@dataclass(frozen=True)
class ReportModel:
    key: str
    name: str
    description: str
    icon: str


@dataclass(frozen=True)
class ReportDocument:
    model: ReportModel
    source: str
    scope: str
    period: str
    records: int
    indicators: tuple[tuple[str, Any], ...]
    tables: tuple[tuple[str, pd.DataFrame], ...]
    alerts: tuple[str, ...]
    methodology: tuple[str, ...]


REPORT_MODELS = (
    ReportModel("executive", "Relatório Executivo", "Síntese gerencial com KPIs, evolução, produtos, lojas e alertas.", "◇"),
    ReportModel("sales", "Relatório de Vendas", "Receita, volume, evolução e comparação temporal verificável.", "↗"),
    ReportModel("products", "Relatório de Produtos", "Ranking, participação e concentração do portfólio.", "▦"),
    ReportModel("stores", "Relatório de Lojas", "Desempenho, participação e ranking das unidades.", "⌂"),
    ReportModel("quality", "Relatório de Qualidade dos Dados", "Completude, verificações, achados, extremos e limitações.", "✓"),
)
MODEL_BY_KEY = {model.key: model for model in REPORT_MODELS}
FORMATS = {"CSV (ZIP)": ("zip", "application/zip"), "Excel": ("xlsx", "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"), "PDF": ("pdf", "application/pdf")}
EXCEL_SAFE_INTEGER_DIGITS = 15


def _period(df: pd.DataFrame) -> str:
    if "Data" not in df.columns or df.empty:
        return "Não disponível"
    dates = pd.to_datetime(df["Data"], errors="coerce").dropna()
    return "Não disponível" if dates.empty else f"{dates.min():%d/%m/%Y}–{dates.max():%d/%m/%Y}"


def _exact_sum(values: pd.Series) -> Decimal:
    return exact_decimal_sum(values)


def _mean(total: Any, denominator: int) -> Any:
    return total / denominator if denominator else Decimal("0")


def _ranking(df: pd.DataFrame, dimension: str) -> pd.DataFrame:
    if df.empty or dimension not in df.columns:
        return pd.DataFrame(columns=[dimension, "Receita", "Quantidade", "Participação receita (%)", "Participação quantidade (%)"])
    grouped = df.groupby(dimension, as_index=False).agg(
        Receita=("Faturamento", _exact_sum), Quantidade=("Quantidade", "sum")
    )
    revenue_total = sum(grouped["Receita"], Decimal("0"))
    quantity_total = float(grouped["Quantidade"].sum())
    grouped["Participação receita (%)"] = (
        grouped["Receita"].map(lambda value: float((value / revenue_total) * 100))
        if revenue_total
        else 0.0
    )
    grouped["Participação quantidade (%)"] = grouped["Quantidade"].div(quantity_total).mul(100) if quantity_total else 0.0
    return grouped.sort_values(["Receita", dimension], ascending=[False, True]).reset_index(drop=True)


def _concentration(ranking: pd.DataFrame) -> pd.DataFrame:
    shares = ranking.get("Participação receita (%)", pd.Series(dtype=float))
    return pd.DataFrame(
        {
            "Faixa": ["Top 1", "Top 3", "Top 5"],
            "Participação da receita (%)": [float(shares.head(limit).sum()) for limit in (1, 3, 5)],
        }
    )


def _without_attrs(frame: pd.DataFrame) -> pd.DataFrame:
    clean = frame.copy()
    clean.attrs = {}
    return clean


def build_report_document(model_key: str, df: pd.DataFrame, source: str, scope: str) -> ReportDocument:
    if model_key not in MODEL_BY_KEY:
        raise ValueError("Modelo de relatório desconhecido.")
    frame = df.copy(deep=True)
    frame.attrs = df.attrs.copy()
    model = MODEL_BY_KEY[model_key]
    analytics = construir_analytics_state(frame) if not frame.empty else {}
    revenue = _exact_sum(frame.get("Faturamento", pd.Series(dtype=float)))
    quantity = float(pd.to_numeric(frame.get("Quantidade", pd.Series(dtype=float)), errors="coerce").sum())
    products = _ranking(frame, "Produto")
    stores = _ranking(frame, "Loja")
    monthly = pd.DataFrame(columns=["Período", "Receita", "Quantidade"])
    if not frame.empty and "Data" in frame.columns:
        work = frame.assign(Período=pd.to_datetime(frame["Data"]).dt.to_period("M").dt.to_timestamp("M"))
        monthly = work.groupby("Período", as_index=False).agg(
            Receita=("Faturamento", _exact_sum), Quantidade=("Quantidade", "sum")
        )
    alerts: list[str] = []
    methodology = [
        "Receita é a soma de Faturamento no escopo selecionado.",
        "Receita média por registro é a receita total dividida pelo número de registros; denominador zero resulta em 0.",
        "Participações usam o total do mesmo escopo; nenhum registro extremo é removido.",
    ]
    audit = audit_data_quality(frame)
    extreme_count = sum(check.occurrences for check in audit.checks if check.name == "Valores monetários extremos")
    if extreme_count:
        alerts.append(f"{extreme_count} valores monetários extremos foram sinalizados e permanecem nos cálculos.")

    common = (
        ("Receita total", revenue),
        ("Quantidade vendida", quantity),
        ("Registros analisados", len(frame)),
        ("Receita média por registro", _mean(revenue, len(frame))),
    )
    if model_key == "executive":
        tables = (("Evolução temporal", monthly), ("Principais produtos", products.head(10)), ("Principais lojas", stores.head(10)), ("Dados do escopo", _without_attrs(frame)))
        indicators = common + (("Produtos distintos", int(frame["Produto"].nunique()) if "Produto" in frame else 0), ("Lojas distintas", int(frame["Loja"].nunique()) if "Loja" in frame else 0))
    elif model_key == "sales":
        comparison = pd.DataFrame([analytics.get("comparison", {})]) if analytics else pd.DataFrame()
        tables = (("Evolução temporal", monthly), ("Comparação temporal", comparison), ("Ranking de produtos", products), ("Ranking de lojas", stores), ("Dados do escopo", _without_attrs(frame)))
        indicators = common
        methodology.append(str(analytics.get("comparison", {}).get("label", "Comparação temporal indisponível.")))
    elif model_key == "products":
        indicators = (
            ("Produtos distintos", len(products)), ("Receita total", revenue), ("Quantidade vendida", quantity),
            ("Receita média por produto", _mean(revenue, len(products))),
        )
        tables = (("Ranking por produto", products), ("Concentração", _concentration(products)), ("Dados do escopo", _without_attrs(frame)))
    elif model_key == "stores":
        indicators = (
            ("Lojas distintas", len(stores)), ("Receita total", revenue), ("Quantidade vendida", quantity),
            ("Receita média por loja", _mean(revenue, len(stores))),
        )
        comparison = pd.DataFrame([analytics.get("comparison", {})]) if analytics else pd.DataFrame()
        tables = (("Ranking por loja", stores), ("Comparação temporal", comparison), ("Dados do escopo", _without_attrs(frame)))
    else:
        checks = pd.DataFrame([{"Verificação": item.name, "Estado": item.status, "Ocorrências": item.occurrences, "Registros afetados": item.affected_records, "Critério": item.criterion} for item in audit.checks])
        indicators = (
            ("Registros auditados", audit.records_analyzed), ("Completude das células elegíveis (%)", audit.completeness),
            ("Registros únicos com alertas", audit.alerted_records), ("Verificações executadas", audit.executed_checks),
            ("Verificações aprovadas", audit.approved_checks),
        )
        tables = (("Completude", audit.completeness_table), ("Verificações", checks), ("Achados", findings_dataframe(audit)))
        alerts.extend(audit.limitations)
        methodology = ["A auditoria reutiliza o motor de qualidade da aplicação.", *audit.limitations]

    return ReportDocument(model, source, scope, _period(frame), len(frame), tuple(indicators), tuple(tables), tuple(alerts), tuple(methodology))


def _safe_filename(value: str) -> str:
    normalized = unicodedata.normalize("NFKD", value)
    ascii_text = "".join(char for char in normalized if not unicodedata.combining(char))
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", ascii_text).strip("._-")
    return cleaned[:80] or "relatorio"


def _safe_sheet_name(value: str) -> str:
    return re.sub(r"[:\\/?*\[\]]", "_", str(value)).strip()[:31] or "Tabela"


def report_filename(document: ReportDocument, extension: str, generated_at: datetime | None = None) -> str:
    generated_at = generated_at or datetime.now()
    return f"DataFlow_{_safe_filename(document.model.name)}_{generated_at:%Y%m%d_%H%M%S}.{extension}"


def _exact_spreadsheet_value(value: Any) -> Any:
    value = neutralize_excel_formula(value)
    if isinstance(value, bool) or value is None:
        return value
    if isinstance(value, (int, float, Decimal)) and not pd.isna(value):
        try:
            decimal = Decimal(str(value))
        except InvalidOperation:
            return value
        significant_digits = len(decimal.as_tuple().digits)
        integer_digits = max(decimal.adjusted() + 1, 0)
        if significant_digits > EXCEL_SAFE_INTEGER_DIGITS or integer_digits > EXCEL_SAFE_INTEGER_DIGITS:
            return format(decimal, "f")
    return value


def _safe_frame(frame: pd.DataFrame) -> pd.DataFrame:
    safe = frame.copy()
    safe.columns = [neutralize_excel_formula(str(column)) for column in safe.columns]
    for column in safe.columns:
        safe[column] = safe[column].map(_exact_spreadsheet_value)
    return safe


def export_report_csv_zip(document: ReportDocument) -> bytes:
    if not document.tables and not document.indicators:
        raise ValueError("O relatório não possui conteúdo exportável.")
    output = BytesIO()
    with zipfile.ZipFile(output, "w", zipfile.ZIP_DEFLATED) as archive:
        metadata = pd.DataFrame(
            [("Relatório", document.model.name), ("Fonte", document.source), ("Escopo", document.scope), ("Período", document.period), ("Registros", document.records)],
            columns=["Campo", "Valor"],
        )
        tables = [("00_metadados", metadata), ("01_indicadores", pd.DataFrame(document.indicators, columns=["Indicador", "Valor"])), *document.tables]
        for index, (name, table) in enumerate(tables, start=1):
            safe = _safe_frame(table)
            buffer = StringIO()
            safe.to_csv(buffer, index=False, sep=";", lineterminator="\n")
            filename = f"{index:02d}_{_safe_filename(name).casefold()}.csv"
            archive.writestr(filename, "\ufeff" + buffer.getvalue())
        archive.writestr("LEIA-ME.txt", "Arquivos CSV em UTF-8 com BOM e separador ponto e vírgula. Textos iniciados por =, +, - ou @ são neutralizados. Valores acima da precisão segura do Excel são exportados como texto decimal exato.")
    return output.getvalue()


def export_report_xlsx(document: ReportDocument) -> bytes:
    output = BytesIO()
    with pd.ExcelWriter(output, engine="openpyxl") as writer:
        metadata = pd.DataFrame(
            [("Relatório", document.model.name), ("Fonte", document.source), ("Escopo", document.scope), ("Período", document.period), ("Registros", document.records)],
            columns=["Campo", "Valor"],
        )
        _safe_frame(metadata).to_excel(writer, sheet_name="Resumo", index=False, startrow=2)
        _safe_frame(pd.DataFrame(document.indicators, columns=["Indicador", "Valor"])).to_excel(writer, sheet_name="Indicadores", index=False, startrow=2)
        for position, (name, table) in enumerate(document.tables, start=1):
            sheet = _safe_sheet_name(name)
            if sheet in writer.book.sheetnames:
                sheet = f"{sheet[:26]} {position}"
            _safe_frame(table).to_excel(writer, sheet_name=sheet, index=False, startrow=2)
        notes = pd.DataFrame(
            {
                "Alertas": pd.Series(list(document.alerts) or ["Sem alertas adicionais."]),
                "Metodologia": pd.Series(document.methodology),
            }
        )
        _safe_frame(notes).to_excel(
            writer, sheet_name="Metodologia", index=False, startrow=2
        )

    output.seek(0)
    workbook = load_workbook(output)
    for worksheet in workbook.worksheets:
        worksheet["A1"] = f"DATAFLOW FUSION · {document.model.name.upper()}"
        worksheet["A1"].font = Font(size=16, bold=True, color="F8FAFC")
        worksheet["A1"].fill = PatternFill("solid", fgColor="0F172A")
        worksheet.merge_cells(start_row=1, start_column=1, end_row=1, end_column=max(worksheet.max_column, 2))
        for cell in worksheet[3]:
            cell.font = Font(bold=True, color="F8FAFC")
            cell.fill = PatternFill("solid", fgColor="172554")
            cell.alignment = Alignment(horizontal="center")
        worksheet.freeze_panes = "A4"
        worksheet.sheet_view.showGridLines = False
        for column in range(1, worksheet.max_column + 1):
            values = [str(worksheet.cell(row, column).value or "") for row in range(1, min(worksheet.max_row, 100) + 1)]
            worksheet.column_dimensions[get_column_letter(column)].width = min(max(map(len, values), default=8) + 3, 46)
    workbook["Resumo"]["A9"] = "Nota de precisão"
    workbook["Resumo"]["B9"] = "Valores com mais de 15 dígitos significativos são armazenados como texto decimal exato."
    final = BytesIO()
    workbook.save(final)
    return final.getvalue()


def _pdf_value(value: Any) -> str:
    value = neutralize_excel_formula(value)
    if isinstance(value, (pd.Timestamp, datetime)):
        return value.strftime("%d/%m/%Y")
    if isinstance(value, float):
        if abs(value) >= 1_000_000:
            return format(Decimal(str(value)), "f")
        return f"{value:,.2f}".replace(",", "X").replace(".", ",").replace("X", ".")
    if isinstance(value, Decimal):
        return format(value, "f")
    return str(value)


def export_report_pdf(document: ReportDocument) -> bytes:
    output = BytesIO()
    styles = getSampleStyleSheet()
    title = ParagraphStyle("DataFlowTitle", parent=styles["Title"], textColor=colors.HexColor("#132A5E"), alignment=TA_CENTER, fontSize=20, leading=24)
    heading = ParagraphStyle("DataFlowHeading", parent=styles["Heading2"], textColor=colors.HexColor("#173B7A"), spaceBefore=10, spaceAfter=6)
    body = ParagraphStyle("DataFlowBody", parent=styles["BodyText"], fontSize=8.5, leading=11)
    table_header = ParagraphStyle(
        "DataFlowTableHeader",
        parent=body,
        textColor=colors.white,
        fontName="Helvetica-Bold",
    )
    doc = SimpleDocTemplate(output, pagesize=landscape(A4), leftMargin=12 * mm, rightMargin=12 * mm, topMargin=12 * mm, bottomMargin=14 * mm, title=document.model.name, author="DataFlow Fusion Ultimate 2.0")
    story: list[Any] = [Paragraph("DataFlow Fusion Ultimate 2.0", title), Paragraph(document.model.name, heading), Paragraph(f"Fonte: {document.source} · Escopo: {document.scope} · Período: {document.period} · Registros: {document.records}", body), Spacer(1, 6)]
    indicator_rows = [["Indicador", "Valor"], *[[name, _pdf_value(value)] for name, value in document.indicators]]
    indicator_table = Table(indicator_rows, colWidths=[92 * mm, 92 * mm], repeatRows=1)
    indicator_table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#172554")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.35, colors.HexColor("#9FB6D9")), ("FONTNAME", (0, 0), (-1, 0), "Helvetica-Bold"), ("FONTSIZE", (0, 0), (-1, -1), 8), ("ROWBACKGROUNDS", (0, 1), (-1, -1), [colors.white, colors.HexColor("#EFF6FF")])]))
    story.extend([indicator_table, Spacer(1, 8)])
    for name, frame in document.tables:
        if frame.empty:
            continue
        story.append(Paragraph(name, heading))
        preview = frame.head(40).copy()
        preview.attrs = {}
        rows = [[Paragraph(str(neutralize_excel_formula(column)), table_header) for column in preview.columns]]
        rows.extend([[Paragraph(_pdf_value(value), body) for value in row] for row in preview.itertuples(index=False, name=None)])
        width = 260 * mm / max(len(preview.columns), 1)
        table = Table(rows, colWidths=[width] * len(preview.columns), repeatRows=1)
        table.setStyle(TableStyle([("BACKGROUND", (0, 0), (-1, 0), colors.HexColor("#172554")), ("TEXTCOLOR", (0, 0), (-1, 0), colors.white), ("GRID", (0, 0), (-1, -1), 0.25, colors.HexColor("#AFC4E4")), ("FONTSIZE", (0, 0), (-1, -1), 6.5), ("VALIGN", (0, 0), (-1, -1), "TOP")]))
        story.append(table)
        if len(frame) > len(preview):
            story.append(Paragraph(f"Prévia limitada a {len(preview)} de {len(frame)} linhas nesta seção.", body))
        story.append(PageBreak())
    story.append(Paragraph("Alertas e metodologia", heading))
    for text in [*document.alerts, *document.methodology]:
        story.append(Paragraph(f"• {text}", body))

    def footer(canvas, pdf_doc):
        canvas.saveState()
        canvas.setFont("Helvetica", 7)
        canvas.setFillColor(colors.HexColor("#526887"))
        canvas.drawString(12 * mm, 7 * mm, f"Gerado localmente em {datetime.now():%d/%m/%Y %H:%M}")
        canvas.drawRightString(285 * mm, 7 * mm, f"Página {pdf_doc.page}")
        canvas.restoreState()

    doc.build(story, onFirstPage=footer, onLaterPages=footer)
    return output.getvalue()


def export_report(document: ReportDocument, format_label: str) -> tuple[bytes, str, str]:
    if format_label not in FORMATS:
        raise ValueError("Formato de relatório desconhecido.")
    extension, mime = FORMATS[format_label]
    exporters = {"zip": export_report_csv_zip, "xlsx": export_report_xlsx, "pdf": export_report_pdf}
    content = exporters[extension](document)
    if not content:
        raise ValueError("O arquivo gerado está vazio.")
    return content, report_filename(document, extension), mime
