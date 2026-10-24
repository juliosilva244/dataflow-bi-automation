"""Public acceptance checks using synthetic data only, never the operational workbook."""

from decimal import Decimal
from io import BytesIO
from pathlib import Path

import groq
import pytest
from openpyxl import load_workbook
from pypdf import PdfReader
from streamlit.testing.v1 import AppTest

from services.loader import carregar_dados
from services.metrics_engine import construir_analytics_state, exact_decimal_sum
from services.reporting import (
    build_report_document,
    export_report_csv_zip,
    export_report_pdf,
    export_report_xlsx,
)

ROOT = Path(__file__).resolve().parents[1]


class DemoUpload(BytesIO):
    name = "vendas_demo.csv"


@pytest.fixture
def demo_frame():
    frame, _, error = carregar_dados(DemoUpload((ROOT / "data/demo/vendas_demo.csv").read_bytes()))
    assert error is None
    assert len(frame) == 720
    return frame


@pytest.fixture
def no_external_client(monkeypatch):
    def forbidden(*args, **kwargs):
        raise AssertionError("External provider calls are forbidden in public acceptance tests")

    monkeypatch.setattr(groq, "Groq", forbidden)


@pytest.mark.parametrize(
    ("page", "navigation"),
    [
        ("Visão Geral", "overview"),
        ("Análise de Vendas", "sales"),
        ("Produtos", "products"),
        ("Lojas", "stores"),
        ("Inteligência Artificial", "ai"),
        ("Qualidade dos Dados", "quality"),
        ("Relatórios", "reports"),
        ("Configurações", "settings"),
    ],
)
@pytest.mark.parametrize("entrypoint", ["app.py", "main.py"])
def test_public_demo_routes(
    no_external_client, monkeypatch, tmp_path, page, navigation, entrypoint,
):
    from services import loader

    monkeypatch.setattr(
        loader, "DEFAULT_DATA_PATH", tmp_path / "absent-operational.xlsx",
    )
    app = AppTest.from_file(ROOT / entrypoint).run(timeout=30)
    assert not app.exception
    assert not app.error
    assert app.session_state["performance_telemetry"].latest["normalizacao"]["rows"] == 720
    app.button(key=f"nav_{navigation}").click().run(timeout=30)
    assert not app.exception
    assert not app.error
    displayed = [item.value for item in app.markdown]
    displayed.extend(item.proto.body for item in app.get("html"))
    assert any("Demonstração — dados fictícios" in text for text in displayed)
    assert app.session_state["active_page"] == page
    assert app.session_state["ai_enabled"] is False
    assert not app.session_state["groq_api_key"]


def test_public_reports_and_exact_demo_totals(demo_frame):
    import zipfile

    total = exact_decimal_sum(demo_frame["Faturamento"])
    state = construir_analytics_state(demo_frame)
    assert state["faturamento_total_exato"] == total
    assert state["ticket_medio_exato"] == total / Decimal(state["total_pedidos"])
    document = build_report_document("executive", demo_frame, "Demo sintética", "Completo")
    with zipfile.ZipFile(BytesIO(export_report_csv_zip(document))) as archive:
        assert archive.testzip() is None
        assert archive.namelist()
    workbook = load_workbook(BytesIO(export_report_xlsx(document)), data_only=False)
    assert workbook.sheetnames
    assert not any(cell.data_type == "f" for sheet in workbook for row in sheet for cell in row)
    workbook.close()
    pdf = PdfReader(BytesIO(export_report_pdf(document)))
    assert len(pdf.pages) > 0
    assert any((page.extract_text() or "").strip() for page in pdf.pages)


def test_public_ai_state_survives_navigation_without_external_calls(
    demo_frame, no_external_client, monkeypatch
):
    import app as application

    monkeypatch.setattr(
        application, "carregar_dados", lambda **kwargs: (demo_frame.copy(), "Demo sintética", None)
    )
    first = AppTest.from_file(ROOT / "main.py").run(timeout=30)
    first.button(key="nav_settings").click().run(timeout=30)
    first.toggle(key="_settings_ai_enabled").set_value(True).run(timeout=30)
    first.button(key="nav_ai").click().run(timeout=30)
    assert first.session_state["ai_enabled"] is True
    first.button(key="nav_settings").click().run(timeout=30)
    assert first.toggle(key="_settings_ai_enabled").value is True
    second = AppTest.from_file(ROOT / "main.py").run(timeout=30)
    assert second.session_state["ai_enabled"] is False
    assert not second.session_state["groq_api_key"]
    assert not first.exception and not second.exception
