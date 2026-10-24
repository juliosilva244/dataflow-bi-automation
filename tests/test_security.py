from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from services.exporter import gerar_excel
from services.insight_engine import gerar_insights_executivos
from services.metrics_engine import construir_analytics_state
from services.security import escape_html, neutralize_excel_formula


def test_escape_html_protege_dados_nao_confiaveis():
    assert escape_html('<img src=x onerror="alert(1)">') == (
        "&lt;img src=x onerror=&quot;alert(1)&quot;&gt;"
    )


def test_insight_escapa_nome_de_produto():
    insights = gerar_insights_executivos(
        {
            "produto_top": "<script>alert(1)</script>",
            "participacao_produto_top": 50,
        }
    )

    assert "<script>" not in insights[0]["text"]
    assert "&lt;script&gt;" in insights[0]["text"]


def test_neutraliza_prefixos_de_formula_excel():
    for value in ("=1+1", "+cmd", "-2+3", "@SUM(A1:A2)", "  =1+1"):
        assert neutralize_excel_formula(value).startswith("'")


def test_exportacao_excel_nao_grava_formula_maliciosa():
    df = pd.DataFrame(
        {
            "Data": pd.to_datetime(["2024-01-01"]),
            "Loja": ["=HYPERLINK(\"https://example.invalid\")"],
            "Produto": ["@SUM(A1:A2)"],
            "Quantidade": [1],
            "Faturamento": [100],
        }
    )
    state = construir_analytics_state(df)

    output = gerar_excel(df, state)
    workbook = load_workbook(BytesIO(output.getvalue()), data_only=False)
    sheet = workbook["Dados_Filtrados"]

    assert sheet["B5"].value.startswith("'")
    assert sheet["C5"].value.startswith("'")
    assert sheet["B5"].data_type != "f"
    assert "Qualidade_Dados" in workbook.sheetnames

