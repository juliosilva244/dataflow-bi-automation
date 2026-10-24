from io import BytesIO

import pandas as pd
from openpyxl import load_workbook

from components.tables import paginate_dataframe
from services.exporter import gerar_excel
from services.loader import carregar_dados
from services.metrics_engine import construir_analytics_state
from services.performance import (
    SessionScopedCache,
    TelemetryCollector,
    content_signature,
    dataframe_fingerprint,
    get_cached_analytics,
    get_cached_export,
)


def sample_dataframe(rows=12):
    df = pd.DataFrame(
        {
            "Data": pd.date_range("2025-01-01", periods=rows),
            "Loja": [f"Loja {index % 3}" for index in range(rows)],
            "Produto": [f"Produto {index % 4}" for index in range(rows)],
            "Pedido": [f"P-{index}" for index in range(rows)],
            "Quantidade": [(index % 5) + 1 for index in range(rows)],
            "Faturamento": [100.0 + index for index in range(rows)],
        }
    )
    return df


def test_content_signature_is_stable_and_sensitive():
    assert content_signature(b"same") == content_signature(b"same")
    assert content_signature(b"same") != content_signature(b"changed")


def test_session_cache_is_isolated_and_lru_bounded():
    first = SessionScopedCache(max_entries=2)
    second = SessionScopedCache(max_entries=2)
    first.set("a", 1)
    first.set("b", 2)
    first.set("c", 3)
    assert first.get("a") is None
    assert first.get("c") == 3
    assert second.get("c") is None


def test_dataframe_fingerprint_invalidates_on_data_or_filter_change():
    df = sample_dataframe()
    assert dataframe_fingerprint(df) != dataframe_fingerprint(df.iloc[:-1])
    changed = df.copy()
    changed.loc[0, "Faturamento"] += 1
    assert dataframe_fingerprint(df) != dataframe_fingerprint(changed)


def test_analytics_cache_preserves_kpis_and_reports_hit():
    df = sample_dataframe()
    cache = SessionScopedCache()
    telemetry = TelemetryCollector()
    direct = construir_analytics_state(df)
    first, key, first_hit = get_cached_analytics(
        df, cache, telemetry, construir_analytics_state
    )
    second, second_key, second_hit = get_cached_analytics(
        df, cache, telemetry, construir_analytics_state
    )
    for field in ("faturamento_total", "ticket_medio", "total_pedidos", "quantidade_total"):
        assert first[field] == direct[field] == second[field]
    assert key == second_key
    assert first_hit is False
    assert second_hit is True
    assert telemetry.latest["agregacao"]["cache_hit"] is True


def test_export_cache_preserves_workbook_and_formats():
    df = sample_dataframe()
    state = construir_analytics_state(df)
    cache = SessionScopedCache()
    telemetry = TelemetryCollector()
    first, key, first_hit = get_cached_export(
        df, state, "analytics:test", cache, telemetry, gerar_excel
    )
    second, second_key, second_hit = get_cached_export(
        df, state, "analytics:test", cache, telemetry, gerar_excel
    )
    assert first == second
    assert key == second_key
    assert first_hit is False
    assert second_hit is True
    workbook = load_workbook(BytesIO(first), read_only=False, data_only=False)
    data_sheet = workbook["Dados_Filtrados"]
    headers = {data_sheet.cell(4, column).value: column for column in range(1, data_sheet.max_column + 1)}
    assert data_sheet.cell(5, headers["Faturamento"]).number_format == "R$ #,##0.00"
    assert data_sheet.cell(5, headers["Quantidade"]).number_format == "#,##0"
    assert data_sheet.cell(5, headers["Data"]).number_format == "dd/mm/yyyy"
    workbook.close()


def test_pagination_changes_only_the_view():
    df = sample_dataframe(12)
    page, total_pages, normalized_page = paginate_dataframe(df, 2, 5)
    assert page.index.tolist() == [5, 6, 7, 8, 9]
    assert total_pages == 3
    assert normalized_page == 2
    assert len(df) == 12


def test_telemetry_does_not_record_payload_values():
    telemetry = TelemetryCollector()
    with telemetry.measure("teste", rows=3, cache_hit=True):
        pass
    assert set(telemetry.latest["teste"]) == {
        "duration_ms",
        "rows",
        "cache_hit",
        "failed",
    }


def test_file_preparation_cache_uses_content_signature(monkeypatch):
    content = (
        b"Data,Loja,Produto,Quantidade,Faturamento\n"
        b"2025-01-01,Centro,Mouse,2,100.00\n"
    )

    class Upload(BytesIO):
        name = "vendas.csv"

    cache = SessionScopedCache()
    telemetry = TelemetryCollector()
    first, _, first_error = carregar_dados(Upload(content), cache=cache, telemetry=telemetry)
    second, _, second_error = carregar_dados(Upload(content), cache=cache, telemetry=telemetry)
    assert first_error is None
    assert second_error is None
    pd.testing.assert_frame_equal(first, second)
    assert telemetry.latest["carregamento"]["cache_hit"] is True
